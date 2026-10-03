"""Excel exports. Every human text cell goes through the no hyphen rule; Shopify identifiers stay exact."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.common import EXACT_COLUMNS, no_hyphen

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _is_exact(col) -> bool:
    return str(col).strip().lower().replace(" ", "_") in EXACT_COLUMNS


def clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    def header(c):
        if _is_exact(c):
            return "SKU" if str(c).lower() == "sku" else str(c)
        text = no_hyphen(str(c).replace("_", " "))
        return text[:1].upper() + text[1:]

    out.columns = [header(c) for c in out.columns]
    for c in out.columns:
        s = out[c]
        if _is_exact(c):
            continue
        if s.dtype == object:
            out[c] = s.map(lambda v: no_hyphen(v) if isinstance(v, str) else v)
        elif np.issubdtype(s.dtype, np.floating):
            out[c] = s.replace([np.inf, -np.inf], np.nan)
    return out


def write_workbook(path: Path, sheets: dict, summary_lines: list[str] | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        if summary_lines:
            pd.DataFrame({"Summary": [no_hyphen(s) for s in summary_lines]}).to_excel(xw, sheet_name="Read me", index=False)
        for name, df in sheets.items():
            if df is None:
                continue
            sheet = no_hyphen(str(name))[:31] or "Sheet"
            frame = clean_frame(df) if len(df.columns) else pd.DataFrame({"Note": ["Nothing to report"]})
            if frame.empty:
                frame = pd.DataFrame({"Note": ["Nothing to report"]})
            frame.to_excel(xw, sheet_name=sheet, index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"
            for idx, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
                width = 10
                for cell in col:
                    v = cell.value
                    if isinstance(v, (dt.date, dt.datetime)):
                        cell.number_format = "dd mmm yyyy"
                        width = max(width, 12)
                    elif isinstance(v, float):
                        cell.number_format = "0.00" if abs(v) < 1000 else "#,##0"
                        width = max(width, 10)
                    elif v is not None:
                        width = max(width, min(len(str(v)) + 2, 60))
                ws.column_dimensions[get_column_letter(idx)].width = width
                if width >= 60:
                    for cell in col[1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return path


def text_cells(path: Path) -> list[tuple[str, str, str]]:
    """All text cells outside identifier columns, used by the no hyphen acceptance test."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        exact = {i for i, h in enumerate(header) if h is not None and _is_exact(h)}
        out.append((ws.title, "sheet", ws.title))
        for row in rows:
            for i, v in enumerate(row):
                if i not in exact and isinstance(v, str):
                    out.append((ws.title, str(header[i]) if i < len(header) else "", v))
    return out
